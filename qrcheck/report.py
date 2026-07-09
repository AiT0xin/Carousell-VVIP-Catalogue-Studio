"""Stage 4 — emit the XLSX report.

One row per QR/merchant plus coverage rows. Flagged rows are highlighted and a
reason_code column carries the machine-readable outcome. A summary sheet gives
the at-a-glance counts so the operator only opens the detail for flags.
"""
from __future__ import annotations

from collections import Counter
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import ReportRow, ReasonCode

_HEADERS = [
    "page", "printed_handle", "decoded_url", "resolves_to",
    "qr_status", "match_status", "in_master", "reason_code", "detail",
    "evidence",
]

# Reason -> fill colour. PASS/CTA_OK greenish; flags amber/red by severity.
_FILLS = {
    ReasonCode.PASS: "E8F5E9",
    ReasonCode.CTA_OK: "F1F8E9",
    ReasonCode.MISSING_FROM_SHEET: "FFF8E1",
    ReasonCode.NOT_IN_MASTER: "FFF8E1",
    ReasonCode.HANDLE_TYPO: "FFE0B2",
    ReasonCode.RENAMED_HANDLE: "FFCDD2",
    ReasonCode.MISMATCH: "FFCDD2",
    ReasonCode.DEAD_LINK: "FFCDD2",
    ReasonCode.DECODE_FAIL: "FFCDD2",
}
_HEADER_FILL = PatternFill("solid", fgColor="263238")
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _row_values(r: ReportRow) -> list:
    return [
        r.page or "", r.printed_handle, r.decoded_url, r.resolves_to,
        r.qr_status, r.match_status, r.in_master, r.reason_code.value,
        r.detail, r.screenshot_path,
    ]


def write_report(rows: Iterable[ReportRow], out_path: str,
                 flagged_first: bool = True) -> dict:
    rows = list(rows)
    if flagged_first:
        rows.sort(key=lambda r: (not r.is_flag, r.page or 9999, r.reason_code.value))

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    # header
    for ci, h in enumerate(_HEADERS, start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for ri, r in enumerate(rows, start=2):
        fill = PatternFill("solid", fgColor=_FILLS.get(r.reason_code, "FFFFFF"))
        for ci, val in enumerate(_row_values(r), start=1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill = fill
            c.border = _BORDER
            c.alignment = Alignment(vertical="top", wrap_text=(ci in (3, 4, 9)))
        # bold the reason_code for flags
        if r.is_flag:
            ws.cell(row=ri, column=8).font = Font(bold=True, color="B71C1C")

    widths = [6, 26, 42, 26, 12, 28, 10, 18, 40, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # summary sheet
    counts = Counter(r.reason_code for r in rows)
    s = wb.create_sheet("Summary")
    s["A1"] = "reason_code"; s["B1"] = "count"
    s["A1"].font = s["B1"].font = Font(bold=True)
    flagged = sum(v for k, v in counts.items() if k.is_flag)
    s["D1"] = "QRs / rows total"; s["E1"] = len(rows)
    s["D2"] = "flagged (need review)"; s["E2"] = flagged
    s["D2"].font = Font(bold=True, color="B71C1C")
    for i, (code, n) in enumerate(sorted(counts.items(), key=lambda kv: kv[0].value), start=2):
        s.cell(row=i, column=1, value=code.value)
        s.cell(row=i, column=2, value=n)
        s.cell(row=i, column=1).fill = PatternFill("solid", fgColor=_FILLS.get(code, "FFFFFF"))
    for col, w in (("A", 20), ("B", 8), ("D", 22), ("E", 10)):
        s.column_dimensions[col].width = w

    wb.save(out_path)
    return {"total": len(rows), "flagged": flagged, "counts": dict(counts)}
